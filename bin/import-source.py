from docx import Document
from pathlib import Path
from PIL import Image
import json, re, unicodedata
root=Path.cwd()
doc=Document('C:/Users/Elias/Downloads/INFORMACION PARA PAGINA WEB.docx')
records=[]
def slug(t):
 return re.sub(r'[^a-z0-9]+','-',unicodedata.normalize('NFKD',t).encode('ascii','ignore').decode().lower()).strip('-')
def clean(t):
 return re.sub(r'\s+',' ',t).strip()
for i,table in enumerate(doc.tables):
 rows=[[clean(c.text) for c in r.cells] for r in table.rows]
 if i==2: continue
 for row in rows[1:]:
  if not row[0]: continue
  title=row[0]
  line='adultos' if i==0 else 'essentials' if i==1 else 'edge'
  track='' if i<3 else {3:'explorer',4:'builders',5:'maker',6:'engineering'}[i]
  full=row[2] if i==0 else title
  desc=row[3] if i==0 else row[3] if i==1 and len(row)>3 and row[3] else row[2]
  age_min,age_max= {'explorer':(3,7),'builders':(8,11),'maker':(12,14),'engineering':(15,None)}.get(track,(None,None))
  if line=='essentials':
   if 'BABY' in title: age_min,age_max=3,5
   elif 'JUNIOR' in title: age_min,age_max=5,7
   elif 'IMPRESION' in title: age_min,age_max=6,12
  tech='LEGO' if 'LEGO' in title or track=='builders' else 'Arduino' if 'ARDUINO' in title or track=='maker' else 'Diseño 3D' if '3D' in title or 'PCB' in title else 'Drones' if 'DRON' in title else 'Robótica'
  if 'mioplécticos' in desc: desc=desc.replace('sensores mioplécticos','sensores')
  title=title.title().replace(' Bot',' Bot').replace('Arduino','Arduino').replace('Lego','LEGO').replace('3D','3D').replace('3d','3D').replace('Pcb','PCB').replace('Fpv','FPV').replace('Cnc','CNC')
  for a,b in [('Robotica','Robótica'),('Impresion','Impresión'),('Electronica','Electrónica'),('Basica','Básica'),('Basico','Básico'),('Protesis','Prótesis'),('Parvulos','Párvulos'),('Autonomo','Autónomo'),('Dinamicos','Dinámicos'),('Programacion','Programación'),('Domotica','Domótica'),('Linea','Línea')]: title=title.replace(a,b)
  img='robot-hero.webp'
  if track=='builders' or tech=='LEGO': img='robot-builder.webp'
  if track in ['maker','engineering'] or tech=='Arduino' or any(t in slug(title) for t in ['electronica','microsoldadura','domotica','python']): img='electronics.webp'
  if any(t in slug(title) for t in ['electronica','microsoldadura']): tech='Electrónica'
  duration=3 if line!='adultos' or 'PARVULOS' in row[0] else 1
  r=dict(slug=slug(title),title=title,full_title=full,description=desc,line=line,track=track,age_min=age_min,age_max=age_max,technology=tech,duration_months=duration,weekly_hours=3 if line=='edge' else 2 if line=='essentials' else None,schedule='Entre semana' if line=='edge' else 'Sábados' if line=='essentials' else '',price=None,price_period='',materials='',requirements='',syllabus='',start_date='',location='',image=img,status='published',featured=1 if row[0] in ['EXPLORER SAFARI BOT','BUILDER SOCCER BOT','MAKER DOMOTICA BOT'] else 0)
  records.append(r)
(root/'database/courses.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
for name in ['robot-hero','roboticlub','open-games','afiche-2026']:
 im=Image.open(root/f'public/assets/images/{name}.png').convert('RGB')
 im.save(root/f'public/assets/images/{name}.webp','WEBP',quality=90)
print('Programas importados:',len(records))
print('Tablas:',len(doc.tables))
print({line:sum(r['line']==line for r in records) for line in ['adultos','essentials','edge']})
